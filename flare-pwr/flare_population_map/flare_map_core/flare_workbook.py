"""Read NOTBADTRAD plume inputs, source terms, and dose results from FLARE workbooks."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .deposition import ReleaseGroup
from .dispersion import DispersionInputs, DoseCurve


@dataclass(frozen=True)
class FlarePlumeCase:
    dispersion: DispersionInputs
    eab_curve: DoseCurve
    lpz_curve: DoseCurve
    eab_distance_m: float
    lpz_distance_m: float
    gap_duration_hr: float
    early_duration_hr: float
    eab_integration_hr: float
    lpz_integration_hr: float
    release_groups: tuple[ReleaseGroup, ...]
    source_path: Path


def _to_float(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _input_echo_values(workbook) -> dict[str, Any]:
    if "Input Echo" not in workbook.sheetnames:
        return {}
    ws = workbook["Input Echo"]
    values: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            values[str(row[0]).strip()] = row[1]
    return values


def _dose_curves(workbook) -> tuple[DoseCurve, DoseCurve]:
    if "Dose" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain a 'Dose' sheet")
    ws = workbook["Dose"]
    start = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and str(row[0]).strip() == "Distance (m)":
            start = i + 1
            break
    if start is None:
        raise ValueError("Dose sheet does not contain the distance-dose table")
    distances, eab, lpz = [], [], []
    for row in ws.iter_rows(min_row=start, values_only=True):
        if row[0] is None:
            break
        try:
            distances.append(float(row[0]))
            eab.append(float(row[2]))
            lpz.append(float(row[3]))
        except (TypeError, ValueError, IndexError):
            continue
    return (
        DoseCurve(distance_m=distances, centerline_dose_rem=eab, label="EAB-interval TEDE"),
        DoseCurve(distance_m=distances, centerline_dose_rem=lpz, label="LPZ-interval TEDE"),
    )


def _release_groups(workbook) -> tuple[ReleaseGroup, ...]:
    if "Source Term" not in workbook.sheetnames:
        return tuple()
    ws = workbook["Source Term"]
    groups: list[ReleaseGroup] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, "TOTAL"):
            break
        try:
            groups.append(ReleaseGroup(
                key=str(row[1]),
                name=str(row[0]),
                core_inventory_ci=float(row[2] or 0.0),
                gap_fraction=float(row[3] or 0.0),
                early_fraction=float(row[4] or 0.0),
                total_release_fraction=float(row[7] or 0.0),
            ))
        except (TypeError, ValueError, IndexError):
            continue
    return tuple(groups)


def load_flare_output(path: str | Path) -> FlarePlumeCase:
    """Load meteorology, EAB/LPZ distances, source terms, and dose curves."""
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    values = _input_echo_values(wb)
    eab_curve, lpz_curve = _dose_curves(wb)
    gap_duration = _to_float(values.get("nbt_gap_release_duration_hr"), 0.5)
    early_duration = _to_float(values.get("nbt_early_iv_duration_hr"), 1.3)
    duration = max(_to_float(values.get("nbt_eab_integration_hr"), 2.0), gap_duration + early_duration)
    dispersion = DispersionInputs(
        stability_class=str(values.get("nbt_stability_class", "F")),
        wind_speed_m_s=_to_float(values.get("nbt_wind_speed_m_s"), 1.0),
        release_height_m=_to_float(values.get("nbt_release_height_m"), 0.0),
        release_duration_hr=duration,
    )
    return FlarePlumeCase(
        dispersion=dispersion,
        eab_curve=eab_curve,
        lpz_curve=lpz_curve,
        eab_distance_m=_to_float(values.get("nbt_distance_eab_m"), 16093.44),
        lpz_distance_m=_to_float(values.get("nbt_distance_lpz_m"), 16093.44),
        gap_duration_hr=gap_duration,
        early_duration_hr=early_duration,
        eab_integration_hr=_to_float(values.get("nbt_eab_integration_hr"), 2.0),
        lpz_integration_hr=_to_float(values.get("nbt_lpz_integration_hr"), 24.0),
        release_groups=_release_groups(wb),
        source_path=path,
    )
