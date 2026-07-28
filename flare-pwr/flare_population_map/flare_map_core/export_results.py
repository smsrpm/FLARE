"""Write time-evolving consequence results into a copy of a FLARE output workbook."""
from __future__ import annotations

from io import BytesIO
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .consequence import build_time_evolving_field


def build_output_workbook(
    source_path: str | Path,
    case,
    site_lat: float,
    site_lon: float,
    minimum_radius_miles: float,
    radius_miles: float,
    bearing_deg: float,
    scenario: str,
    deposition_group: str,
    pathway_objective: str,
    seed_text: str,
    interval_count: int = 24,
) -> bytes:
    wb = load_workbook(source_path)
    for name in ("Footprint Summary", "Footprint Intervals", "Footprint Spatial"):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("Footprint Summary")
    summary = [
        ("Result", "Time-Evolving Consequence Footprint"),
        ("Scenario", scenario),
        ("Synthetic intervals", interval_count),
        ("EAB integration time (hr)", case.eab_integration_hr),
        ("LPZ integration time (hr)", case.lpz_integration_hr),
        ("Site latitude", site_lat),
        ("Site longitude", site_lon),
        ("Minimum analysis radius (mi)", minimum_radius_miles),
        ("Maximum analysis radius (mi)", radius_miles),
        ("Maximum-population bearing (deg)", bearing_deg),
        ("Pathway optimization criterion", pathway_objective),
        ("Deposition group", deposition_group),
        ("Random-sequence key", seed_text),
        ("Radial bins", 48),
        ("Azimuth bins", 72),
        ("Note", "Synthetic consequence footprint; not a site-specific meteorological reconstruction."),
    ]
    for row in summary:
        ws.append(row)

    # One build supplies the deterministic interval history.
    base = build_time_evolving_field(
        radius_miles, bearing_deg, case.eab_curve, case.lpz_curve, case.dispersion,
        case.eab_integration_hr, case.lpz_integration_hr, interval_count, scenario,
        "Airborne dose", case.release_groups, deposition_group, seed_text,
        interval_count=interval_count, grid_size=80,
    )
    wi = wb.create_sheet("Footprint Intervals")
    wi.append([
        "Step", "Start time (hr)", "End time (hr)", "Bearing (deg)",
        "Wind speed (m/s)", "Stability class", "Incremental interpolation fraction",
        "Cumulative interpolation fraction",
    ])
    wi.append([0, 0.0, case.eab_integration_hr, bearing_deg, case.dispersion.wind_speed_m_s,
               case.dispersion.stability_class, None, 0.0])
    for s in base.steps:
        wi.append([s.step, s.start_hr, s.end_hr, s.bearing_deg, s.wind_speed_m_s,
                   s.stability_class, s.incremental_fraction, s.cumulative_fraction])

    spatial = wb.create_sheet("Footprint Spatial")
    spatial.append([
        "Step", "Elapsed time (hr)", "Radius (m)", "Azimuth (deg)",
        "Cumulative TEDE (rem)", "Cumulative relative deposition",
    ])
    # Precompute each layer once; frame selection below is an array lookup rather
    # than 50 repeated full-field calculations.
    dose = build_time_evolving_field(
        radius_miles, bearing_deg, case.eab_curve, case.lpz_curve, case.dispersion,
        case.eab_integration_hr, case.lpz_integration_hr, 0, scenario,
        "Airborne dose", case.release_groups, deposition_group, seed_text,
        interval_count=interval_count, grid_size=140,
    )
    dep = build_time_evolving_field(
        radius_miles, bearing_deg, case.eab_curve, case.lpz_curve, case.dispersion,
        case.eab_integration_hr, case.lpz_integration_hr, 0, scenario,
        "Ground deposition", case.release_groups, deposition_group, seed_text,
        interval_count=interval_count, grid_size=140,
    )
    radial_values = np.linspace(minimum_radius_miles * 1609.344, radius_miles * 1609.344, 48)
    azimuth_values = np.linspace(0.0, 355.0, 72)
    east_axis = dose.east_m[0, :]
    north_axis = dose.north_m[:, 0]
    for step in range(interval_count + 1):
        dose_field = dose.cumulative_fields[step]
        dep_field = dep.cumulative_fields[step]
        elapsed = case.eab_integration_hr if step == 0 else base.steps[step-1].end_hr
        for r in radial_values:
            for az in azimuth_values:
                br = math.radians(float(az))
                east = r * math.sin(br)
                north = r * math.cos(br)
                ix = int(np.argmin(np.abs(east_axis - east)))
                iy = int(np.argmin(np.abs(north_axis - north)))
                spatial.append([step, elapsed, float(r), float(az),
                                float(dose_field[iy, ix]),
                                float(dep_field[iy, ix])])

    for sheet in (ws, wi, spatial):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
